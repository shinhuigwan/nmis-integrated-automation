from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


@dataclass
class FeePaymentItem:
    excel_row: int
    payment_date: date
    member_name: str
    amount: int
    status: str = "대기"
    detail: str = ""

    @property
    def stable_id(self) -> str:
        return f"fee-{self.excel_row}-{self.payment_date:%Y%m%d}-{self.member_name}-{self.amount}"

    @property
    def month_text(self) -> str:
        return self.payment_date.strftime("%Y-%m")

    @property
    def month_raw(self) -> str:
        return self.payment_date.strftime("%Y%m")

    @property
    def receipt_date_raw(self) -> str:
        return self.payment_date.strftime("%Y%m%d")


@dataclass
class FeePaymentLoadResult:
    items: list[FeePaymentItem]
    skipped_rows: list[str]
    sheet_name: str


def _normalize_header(value: Any) -> str:
    return re.sub(r"[\s·ㆍ&＆]+", "", str(value or "")).lower()


def _parse_excel_date(value: Any, epoch: datetime) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    text = str(value or "").strip()
    digits = re.sub(r"[^0-9]", "", text)
    if len(digits) == 8:
        return datetime.strptime(digits, "%Y%m%d").date()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"날짜 형식을 해석할 수 없습니다: {value}")


def _parse_amount(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError("금액이 숫자가 아닙니다.")
    if isinstance(value, (int, float)):
        amount = int(value)
    else:
        digits = re.sub(r"[^0-9-]", "", str(value or ""))
        if not digits or digits == "-":
            raise ValueError("금액이 비어 있습니다.")
        amount = int(digits)
    if amount <= 0:
        raise ValueError("금액은 0보다 커야 합니다.")
    return amount


def parse_fee_payment_excel(path: str | Path) -> FeePaymentLoadResult:
    """거래내역 시트의 거래일자·성명/상호·금액을 회비 입력 항목으로 읽는다."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"엑셀 파일을 찾을 수 없습니다: {source}")
    # 파일 내용을 메모리에서 열어 openpyxl이 Windows 원본 파일을 계속 점유하지
    # 않게 한다. 자동화 실행 중에도 사용자가 원본 엑셀을 수정/교체할 수 있다.
    workbook = load_workbook(BytesIO(source.read_bytes()), read_only=True, data_only=True)
    try:
        sheet = workbook["거래내역"] if "거래내역" in workbook.sheetnames else workbook.worksheets[0]
        header_row = 0
        column_map: dict[str, int] = {}
        aliases = {
            "date": {"거래일자", "일자", "입금일자"},
            "name": {"성명/상호", "성명상호", "성명", "대표자", "입금자"},
            "amount": {"금액", "입금액", "거래금액"},
        }
        normalized_aliases = {
            key: {_normalize_header(alias) for alias in values}
            for key, values in aliases.items()
        }
        for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            found: dict[str, int] = {}
            for column_index, value in enumerate(row, start=1):
                header = _normalize_header(value)
                for key, names in normalized_aliases.items():
                    if header in names:
                        found[key] = column_index
            if set(found) == {"date", "name", "amount"}:
                header_row = row_index
                column_map = found
                break
        if not header_row:
            raise ValueError("거래일자, 성명/상호, 금액 헤더를 찾지 못했습니다.")

        items: list[FeePaymentItem] = []
        skipped: list[str] = []
        max_column = max(column_map.values())
        for excel_row, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, max_col=max_column, values_only=True),
            start=header_row + 1,
        ):
            date_value = row[column_map["date"] - 1]
            name_value = row[column_map["name"] - 1]
            amount_value = row[column_map["amount"] - 1]
            if date_value in (None, "") and name_value in (None, "") and amount_value in (None, ""):
                continue
            try:
                member_name = str(name_value or "").strip()
                if not member_name:
                    raise ValueError("성명/상호가 비어 있습니다.")
                items.append(
                    FeePaymentItem(
                        excel_row=excel_row,
                        payment_date=_parse_excel_date(date_value, workbook.epoch),
                        member_name=member_name,
                        amount=_parse_amount(amount_value),
                    )
                )
            except Exception as exc:
                skipped.append(f"엑셀 {excel_row}행: {exc}")
        items.sort(key=lambda item: (item.payment_date, item.excel_row))
        return FeePaymentLoadResult(items=items, skipped_rows=skipped, sheet_name=sheet.title)
    finally:
        workbook.close()


def find_monthly_fee_row_index(row_texts: Iterable[str]) -> int | None:
    for index, text in enumerate(row_texts):
        if "월회비" in re.sub(r"\s+", "", str(text or "")):
            return index
    return None


def find_member_monthly_fee_row_index(
    row_texts: Iterable[str],
    member_name: str,
    fee_type_labels: Iterable[str] | None = None,
) -> int | None:
    normalized_member = re.sub(r"\s+", "", str(member_name or ""))
    texts = list(row_texts)
    labels = list(fee_type_labels) if fee_type_labels is not None else None
    for index, text in enumerate(texts):
        normalized_text = re.sub(r"\s+", "", str(text or ""))
        label_is_monthly = (
            re.sub(r"\s+", "", str(labels[index] if labels and index < len(labels) else ""))
            == "월회비"
            if labels is not None
            else "월회비" in normalized_text
        )
        if normalized_member in normalized_text and label_is_monthly:
            return index
    return None


def _fee_page_ready(page: Any) -> bool:
    return (
        page.locator("input[name='stdYearMonth']:visible").count() > 0
        and page.locator("input[name='ceoMemberName']:visible").count() > 0
    )


def ensure_nmis_logged_in(
    page: Any,
    user_id: str,
    password: str,
    log_cb: Callable[[str], None] | None = None,
) -> bool:
    """세션 만료로 로그인 화면이 표시되면 저장된 NMIS 계정으로 재로그인한다."""
    log = log_cb or (lambda _message: None)
    user_input = page.locator("#userId:visible")
    password_input = page.locator("#userPass:visible")
    if user_input.count() == 0 or password_input.count() == 0:
        return False
    if not user_id or not password:
        raise RuntimeError("NMIS 로그인 세션이 만료되었지만 저장된 로그인 계정이 없습니다.")

    log("NMIS 로그인 세션이 만료되어 자동으로 다시 로그인합니다.")
    user_input.first.fill(user_id)
    password_input.first.fill(password)
    password_input.first.press("Enter")

    # 로그인 직후 비밀번호 갱신 안내가 나오면 기존 자동 로그인과 같이 취소한다.
    for _ in range(40):
        cancel_button = page.locator("button[ng-click*='fnCancel']:visible")
        if cancel_button.count() > 0:
            cancel_button.first.click(force=True)
        if page.locator("#userId:visible").count() == 0:
            break
        page.wait_for_timeout(250)

    if page.locator("#userId:visible").count() > 0:
        raise RuntimeError("NMIS 자동 재로그인에 실패했습니다. 저장된 계정 정보를 확인하세요.")

    top_menu = page.locator("#menuKFB_BH:visible")
    if top_menu.count() == 0:
        top_menu = page.get_by_role("link", name="회비", exact=True)
    top_menu.first.wait_for(state="visible", timeout=10000)
    log("NMIS 자동 재로그인 완료")
    return True


def navigate_to_fee_management(page: Any, log_cb: Callable[[str], None] | None = None) -> None:
    """회비 > 회비＆가입비관리 > 회비 및 가입금관리 화면으로 이동한다."""
    log = log_cb or (lambda _message: None)
    if _fee_page_ready(page):
        return
    log("NMIS 메뉴 이동: 회비 > 회비＆가입비관리 > 회비 및 가입금관리")

    # 잠재회원·접수대장 자동화와 동일하게 AngularJS 내부 화면 경로로 바로
    # 이동한다. 슬래시/점 등을 제거했을 때 feecontrollist와 정확히 같은
    # 상태만 허용하므로 CMS·징수현황 화면을 잘못 고르지 않는다.
    try:
        route_result = page.evaluate(
            """() => {
                const targets = [
                    document.querySelector('[ng-app]'),
                    document.body,
                    document.documentElement,
                    document.querySelector('#wrap'),
                    document.querySelector('#container')
                ].filter(Boolean);
                for (const target of targets) {
                    try {
                        const injector = window.angular && angular.element(target).injector();
                        if (!injector || !injector.has('$state')) continue;
                        const state = injector.get('$state');
                        const names = state.get().map(item => item && item.name).filter(Boolean);
                        const normalize = value => String(value).toLowerCase().replace(/[^a-z0-9]/g, '');
                        const route = names.find(name => normalize(name) === 'feecontrollist');
                        if (!route) return {route: '', error: 'exact fee/control/list route not found'};
                        state.go(route);
                        return {route};
                    } catch (_) {}
                }
                return {route: '', error: 'NMIS Angular state service not found'};
            }"""
        )
        if route_result.get("route"):
            log(f"NMIS 내부 경로로 바로 이동합니다: {route_result['route']}")
            page.locator("input[name='stdYearMonth']:visible").first.wait_for(
                state="visible", timeout=15000
            )
            page.locator("input[name='ceoMemberName']:visible").first.wait_for(
                state="visible", timeout=15000
            )
            if _fee_page_ready(page):
                log("NMIS 회비 및 가입금관리 화면 이동 완료")
                return
    except Exception as exc:
        log(f"NMIS 내부 경로 이동 대기 실패 — 메뉴 클릭 방식으로 전환합니다: {exc}")

    # 내부 경로를 제공하지 않는 화면 버전에서는 실제 고유 ID 클릭을 예비 수단으로 사용한다.
    # 상단 메뉴를 누른 뒤 왼쪽 메뉴가 Angular에서 늦게 생성되므로 충분히 기다린다.
    top_menu = page.locator("#menuKFB_BH:visible")
    if top_menu.count() == 0:
        top_menu = page.get_by_role("link", name="회비", exact=True)
    if top_menu.count() == 0:
        raise RuntimeError("NMIS 상단의 '회비' 메뉴를 찾지 못했습니다.")

    last_error: Exception | None = None
    for attempt in range(1, 4):
        if _fee_page_ready(page):
            log("NMIS 회비 및 가입금관리 화면 이동 완료")
            return
        try:
            top_menu.first.click(force=True)
            submenu = page.locator("#feecontrollist")
            submenu.first.wait_for(state="attached", timeout=15000)
            # 메뉴 그룹이 접혀 있더라도 고유 ID에 강제 클릭하면 Angular menuLink가 실행된다.
            submenu.first.click(force=True)
            page.locator("input[name='stdYearMonth']:visible").first.wait_for(
                state="visible", timeout=15000
            )
            page.locator("input[name='ceoMemberName']:visible").first.wait_for(
                state="visible", timeout=15000
            )
            if _fee_page_ready(page):
                log("NMIS 회비 및 가입금관리 화면 이동 완료")
                return
        except Exception as exc:
            last_error = exc
            log(f"회비 및 가입금관리 메뉴 이동 재시도 {attempt}/3")
            page.wait_for_timeout(1000)
            top_menu = page.locator("#menuKFB_BH:visible")
            if top_menu.count() == 0:
                top_menu = page.get_by_role("link", name="회비", exact=True)

    detail = f" ({last_error})" if last_error else ""
    raise RuntimeError(
        "'회비 > 회비＆가입비관리 > 회비 및 가입금관리' 화면으로 이동하지 못했습니다."
        + detail
    )


def _visible_grid_rows_with_amount(page: Any) -> list[Any]:
    amount_inputs = page.locator("input[ng-model*=\"monthInAmt\"]:visible")
    rows: list[Any] = []
    for index in range(amount_inputs.count()):
        row = amount_inputs.nth(index).locator(
            "xpath=ancestor::div[contains(concat(' ', normalize-space(@class), ' '), ' ui-grid-row ')][1]"
        )
        if row.count() > 0:
            rows.append(row)
    return rows


def _fee_type_label_from_row(row: Any) -> str:
    """행의 회비구분 select에서 실제 선택값(가입금/월회비)을 읽는다."""
    selects = row.locator("select:visible")
    for index in range(selects.count()):
        selected_option = selects.nth(index).locator("option:checked")
        if selected_option.count() == 0:
            continue
        label = re.sub(r"\s+", "", str(selected_option.first.inner_text() or ""))
        if label in {"가입금", "월회비"}:
            return label
    return ""


def _can_fill_input(locator: Any) -> bool:
    """readonly/disabled 검색조건은 Playwright fill 대상에서 제외한다."""
    return (
        locator.get_attribute("readonly") is None
        and locator.get_attribute("disabled") is None
        and locator.get_attribute("aria-disabled") != "true"
    )


def _month_value_matches(value: Any, expected_date: date) -> bool:
    return re.sub(r"[^0-9]", "", str(value or "")) == expected_date.strftime("%Y%m")


def register_fee_payment_on_nmis(
    page: Any,
    item: FeePaymentItem,
    log_cb: Callable[[str], None] | None = None,
    user_id: str = "",
    password: str = "",
) -> dict[str, Any]:
    """대표자를 조회해 월회비 행에 입금액·입금일자를 입력하고 저장한다."""
    log = log_cb or (lambda _message: None)
    if user_id and password:
        ensure_nmis_logged_in(page, user_id, password, log)
    navigate_to_fee_management(page, log)

    # 이전 조회 조건이 남아 있으면 대표자 검색 결과가 누락되므로 충돌 필터를 비운다.
    for name in ("memberName", "businessReportNo", "hDongName", "bDongName"):
        locator = page.locator(f"input[name='{name}']:visible")
        if locator.count() > 0 and _can_fill_input(locator.first):
            locator.first.fill("", timeout=3000)

    for name, label in (
        ("memberType", "전체"),
        ("feeType", "전체"),
        ("treatType", "대상자"),
        ("monthFeeType", "전체"),
    ):
        locator = page.locator(f"select[name='{name}']:visible")
        if locator.count() > 0:
            locator.first.select_option(label=label)

    month_input = page.locator("input[name='stdYearMonth']:visible").first
    representative_input = page.locator("input[name='ceoMemberName']:visible").first
    # 이 입력칸은 maxlength=6이므로 2026-01(7자)을 넣으면 2026-0으로
    # 잘린다. 원시 6자리 YYYYMM을 넣으면 NMIS가 화면에서 YYYY-MM으로 표시한다.
    month_input.fill(item.month_raw)
    page.wait_for_timeout(500)
    # 달력 아이콘은 누르지 않는다. 다른 입력박스를 클릭해 blur/change를
    # 발생시킨 뒤 화면에 확정된 기준년월을 다시 읽어 정확성을 검증한다.
    representative_input.click(force=True)
    page.wait_for_timeout(200)
    confirmed_month = month_input.input_value()
    if not _month_value_matches(confirmed_month, item.payment_date):
        raise RuntimeError(
            "기준년월 입력값이 확정되지 않았습니다: "
            f"입력 예정={item.month_text}, 화면값={confirmed_month or '빈 값'}"
        )
    log(f"기준년월 입력 확인 완료: {confirmed_month}")
    representative_input.fill(item.member_name)
    page.wait_for_timeout(500)
    representative_input.press("Tab")
    page.wait_for_timeout(500)

    # 대표자 입력칸의 Enter 이벤트에 의존하지 않고 이 검색폼의 조회 버튼을
    # 명시적으로 한 번 눌러야 그리드 결과가 실제로 갱신된다.
    search_form = month_input.locator("xpath=ancestor::form[1]")
    search_button = search_form.get_by_role("button", name="조회", exact=True)
    if search_button.count() == 0:
        search_button = page.locator(
            "form:has(input[name='stdYearMonth']) button[ng-click*='fnSearch']:visible, "
            "form:has(input[name='stdYearMonth']) button[ng-click*='Search']:visible, "
            "form:has(input[name='stdYearMonth']) button:has-text('조회'):visible"
        )
    if search_button.count() == 0:
        raise RuntimeError("회비 및 가입금관리 검색폼의 '조회' 버튼을 찾지 못했습니다.")
    search_button.first.click(force=True)
    log(
        f"회비 조회 실행: 기준년월={item.month_text}, "
        f"대표자={item.member_name}"
    )
    page.wait_for_timeout(2000)

    rows = _visible_grid_rows_with_amount(page)
    row_texts = [row.inner_text() for row in rows]
    fee_type_labels = [_fee_type_label_from_row(row) for row in rows]
    monthly_index = find_member_monthly_fee_row_index(
        row_texts, item.member_name, fee_type_labels
    )
    if monthly_index is None:
        normalized_member = re.sub(r"\s+", "", item.member_name)
        member_row_texts = [
            text for text in row_texts
            if normalized_member in re.sub(r"\s+", "", text)
        ]
        fee_types = [
            label
            for label in ("가입금", "월회비")
            if any(
                normalized_member in re.sub(r"\s+", "", row_texts[index])
                and fee_type_labels[index] == label
                for index in range(len(row_texts))
            )
        ]
        found_label = ", ".join(fee_types) if fee_types else "조회 결과 없음"
        log(
            f"⚠️ 월회비 행 없음: {item.member_name} / {item.payment_date:%Y-%m-%d} "
            f"({found_label})"
        )
        return {"success": False, "missing_month_fee": True, "found": found_label}

    monthly_row = rows[monthly_index]
    confirmed_fee_type = _fee_type_label_from_row(monthly_row)
    if confirmed_fee_type != "월회비":
        raise RuntimeError(
            f"안전 중단: 선택하려는 행의 회비구분이 월회비가 아닙니다 ({confirmed_fee_type or '확인 불가'})."
        )

    # 금액과 날짜는 반드시 확인된 월회비 행 내부 요소만 사용한다. 선택 열만
    # 왼쪽 고정 그리드에 있으므로 같은 행 순번의 체크박스를 사용한다.
    amount_input = monthly_row.locator("input[ng-model*='monthInAmt']:visible").first
    receipt_date_input = monthly_row.locator("input[ng-model*='receiptDate']:visible").first
    selected_checkboxes = page.locator(
        "input[type='checkbox'][ng-model*='row'][ng-model*='selected']:visible"
    )
    checkbox_count = selected_checkboxes.count()
    if amount_input.count() == 0 or receipt_date_input.count() == 0 or monthly_index >= checkbox_count:
        raise RuntimeError(
            "월회비 행의 입력 요소를 찾지 못했습니다: "
            f"월회비행={monthly_index + 1}, 금액={amount_input.count()}, "
            f"입금일자={receipt_date_input.count()}, 선택={checkbox_count}"
        )
    selected_checkbox = selected_checkboxes.nth(monthly_index)

    amount_input.fill(str(item.amount))
    page.wait_for_timeout(500)
    amount_input.press("Tab")
    page.wait_for_timeout(500)
    receipt_date_input.fill(item.receipt_date_raw)
    page.wait_for_timeout(500)
    receipt_date_input.press("Tab")
    page.wait_for_timeout(500)
    selected_checkbox.check(force=True)
    page.wait_for_timeout(500)

    if amount_input.input_value().replace(",", "") != str(item.amount):
        raise RuntimeError("당월입금액 입력값을 확인하지 못했습니다.")
    if re.sub(r"[^0-9]", "", receipt_date_input.input_value()) != item.receipt_date_raw:
        raise RuntimeError("입금일자 입력값을 확인하지 못했습니다.")
    if not selected_checkbox.is_checked():
        raise RuntimeError("월회비 행 선택 체크가 적용되지 않았습니다.")

    save_button = page.locator(
        "button:has(span.button_icon[lang-code='save']):visible, "
        "button[ng-click*='fnSave']:visible"
    )
    if save_button.count() == 0:
        raise RuntimeError("회비 및 가입금관리 저장 버튼을 찾지 못했습니다.")
    save_button.first.click(force=True)
    page.wait_for_timeout(500)

    confirm_button = page.locator("button.btn-success[ng-click*='fnConfirm']:visible")
    confirm_button.first.wait_for(state="visible", timeout=5000)
    confirm_button.first.click(force=True)
    page.wait_for_timeout(500)

    saved_button = page.locator("button.btn-success[ng-click*='fnClose']:visible")
    saved_button.first.wait_for(state="visible", timeout=7000)
    saved_button.first.click(force=True)
    page.wait_for_timeout(500)
    log(
        f"✅ 월회비 저장 완료: {item.member_name} / {item.payment_date:%Y-%m-%d} / {item.amount:,}원"
    )
    return {"success": True, "missing_month_fee": False}
